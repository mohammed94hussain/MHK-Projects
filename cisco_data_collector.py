import getpass
import time
import re
from netmiko import ConnectHandler
from openpyxl import Workbook, load_workbook
import sys

# --- Prompt for username and password (using getpass for security) ---
print("--- Authentication ---")
username = input("Enter your username: ")
password = getpass.getpass()
print("----------------------")

# Read IP addresses from IP.xlsx
try:
    wb_ip = load_workbook('IP.xlsx')
    sheet_ip = wb_ip.active
    ips = [cell.value for cell in sheet_ip['A'][1:] if cell.value]
except FileNotFoundError:
    print("Error: IP.xlsx not found. Please create it and add IP addresses in column A, starting from row 2.")
    sys.exit(1)

# Create a new workbook for the output
wb_db = Workbook()
sheet_db = wb_db.active
sheet_db.title = "DataBase"

# Write headers to the output file
headers = [
    "Hostname", "IP", "Port", "Description", "Status", "VLAN",
    "Neighbor", "Port Speed", "Link Uptime", "Device Serial",
    "Power Supply Serials", "Transceiver Info", "Port Serial"
]
sheet_db.append(headers)

# Loop through each IP address
for ip in ips:
    print(f"\n🚀 Connecting to {ip}...")
    device = {
        'device_type': 'cisco_ios',
        'host': ip,
        'username': username,
        'password': password,
        'conn_timeout': 60,
        'auth_timeout': 60,
        'global_delay_factor': 2,
    }

    try:
        net_connect = ConnectHandler(**device)
        hostname = net_connect.find_prompt()[:-1]
        print(f"✅ Successfully logged in and prompt found: {hostname}")

        # Get inventory, including transceivers
        show_inventory_raw = net_connect.send_command('show inventory', use_textfsm=False)
        device_serial = "N/A"
        power_supply_serials = []
        transceivers = {}

        # Parse chassis serial
        chassis_match = re.search(r'NAME: "Chassis.*", DESCR: ".*"\nPID: .*, VID: .*, SN: (\S+)', show_inventory_raw)
        if chassis_match:
            device_serial = chassis_match.group(1)

        # Parse power supply serials
        for match in re.finditer(r'NAME: "Power Supply Module \d+", DESCR: ".*"\nPID: .*, VID: .*, SN: (\S+)', show_inventory_raw):
            power_supply_serials.append(match.group(1))

        # Parse transceiver serials and info from inventory
        for match in re.finditer(r'NAME: "(\S+Ethernet\S+)", DESCR: "(\S+)"\nPID: (\S+)\s+, VID: \S+\s+, SN: (\S+)', show_inventory_raw):
            port_name, descr, pid, sn = match.groups()
            if "transceiver" in descr.lower() or "sfp" in descr.lower():
                 transceivers[port_name] = f"Type: {pid}, SN: {sn}"

        # Get interface status (Base for loop)
        show_int_status = net_connect.send_command('show interfaces status', use_textfsm=True)

        # FIX FOR MISSING LLDP NEIGHBOR KEYS
        show_lldp_neighbors = net_connect.send_command('show lldp neighbors detail', use_textfsm=True)
        neighbors = {}
        if isinstance(show_lldp_neighbors, list):
            for neighbor in show_lldp_neighbors:
                local_port = neighbor.get('local_interface') or neighbor.get('local_port')
                if local_port:
                    neighbors[local_port] = neighbor.get('neighbor_system_name') or neighbor.get('neighbor') or 'N/A'

        # FIX FOR KEY ERROR: 'desc' (Interface Descriptions)
        show_int_desc = net_connect.send_command('show interfaces description', use_textfsm=True)
        descriptions = {}
        if isinstance(show_int_desc, list):
            for item in show_int_desc:
                port_key = item.get('port') or item.get('interface')
                desc_value = item.get('desc') or item.get('description', 'N/A')

                if port_key:
                    descriptions[port_key] = desc_value

        # Get interface uptime from raw output
        show_interfaces_raw = net_connect.send_command('show interfaces', use_textfsm=False)
        interface_uptimes = {}
        # Simple regex to find "Last link flapped..."
        for match in re.finditer(r'(\S+ is .*?)\n(?:.|\n)*?Last link flapped (\S+)', show_interfaces_raw):
            interface_name = match.group(1).split()[0]
            uptime = match.group(2)
            interface_uptimes[interface_name] = uptime

        # Process each interface
        for interface in show_int_status:
            # UNIFIED PORT KEY: Use this key for all lookups!
            port = interface.get('port') or interface.get('interface', 'N/A')

            description = descriptions.get(port, 'N/A')
            status = interface.get('status', 'N/A')

            # FIX FOR KEY ERROR: 'vlan'
            vlan = interface.get('vlan') or interface.get('vlan_id') or interface.get('access_vlan', 'N/A')

            neighbor = neighbors.get(port, 'N/A')
            port_speed = interface.get('speed', 'N/A')
            link_uptime = interface_uptimes.get(port, 'N/A')
            transceiver_info = transceivers.get(port, 'N/A')
            port_serial = transceiver_info.split('SN: ')[1] if transceiver_info and 'SN: ' in transceiver_info else 'N/A'

            row_data = [
                hostname,
                ip,
                port,
                description,
                status,
                vlan,
                neighbor,
                port_speed,
                link_uptime,
                device_serial,
                ", ".join(power_supply_serials),
                transceiver_info,
                port_serial
            ]
            sheet_db.append(row_data)

        net_connect.disconnect()
        print(f"Successfully collected data from {ip}")

    except Exception as e:
        error_type = "AUTHENTICATION FAILED" if "Authentication to device failed" in str(
            e) else "COMMAND FAILED (Post-Login)"
        print(f"❌ Failure on {ip}: {e}")

        error_message = f"{error_type}: {str(e).splitlines()[0]}"
        sheet_db.append([error_type, ip, error_message, "", "", "", "", "", "", "", "", ""])

    print("⏳ Waiting for 10 seconds...")
    time.sleep(10)

# Save the workbook
wb_db.save('DataBase.xlsx')

print("\n--- Script finished. Data saved to DataBase.xlsx ---")
